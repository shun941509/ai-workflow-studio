# Operational file helper. Original inputs are read-only; results use a new directory.
import sys, json, os, re, hashlib
from pathlib import Path
def fail(message): raise ValueError(message)
def pairs(values):
    result={}
    for key,value in values:
        if key in result: fail('Duplicate JSON key.')
        result[key]=value
    return result
def text(value): return isinstance(value,str) and bool(value.strip())
def digest(data): return hashlib.sha256(data).hexdigest()
def source(value, maximum=33554432, directory=False):
    if not text(value) or not Path(value).is_absolute(): fail('Approved absolute input path required.')
    candidate=Path(value)
    for part in [candidate,*candidate.parents]:
        if part.is_symlink() or (hasattr(part,'is_junction') and part.is_junction()): fail('Links and junctions are not supported.')
    result=candidate.resolve(strict=True)
    if directory:
        if not result.is_dir(): fail('Input directory required.')
    elif not result.is_file() or result.stat().st_size>maximum: fail('Input file missing or exceeds supported size.')
    return result
def output_directory(value):
    if not text(value) or not Path(value).is_absolute(): fail('New absolute output directory required.')
    candidate=Path(value); source(str(candidate.parent),directory=True)
    try: candidate.mkdir(exist_ok=False)
    except FileExistsError: fail('Output directory already exists.')
    return candidate
def safe_name(value):
    if not text(value) or len(value)>150 or value in ['.','..'] or re.search(r'[<>:"/\\|?*\x00-\x1f]',value) or value.endswith(('.', ' ')) or re.match(r'^(?:con|prn|aux|nul|com[1-9]|lpt[1-9])(?:\.|$)',value,re.I): fail('Unsafe file name.')
    return value
def write_new(path,data):
    with path.open('xb') as target: target.write(data)
def json_bytes(value): return (json.dumps(value,ensure_ascii=False,indent=2,allow_nan=False)+'\n').encode('utf-8')
def run(data):

    # Generate normal WordprocessingML only; no macros, external templates, or remote relationships.
    from docx import Document
    title=data.get('title'); summary=data.get('summary'); sections=data.get('sections'); required=data.get('requiredSections')
    if not text(title) or not isinstance(summary,str) or not isinstance(sections,list) or not sections or not isinstance(required,list) or any(not text(v) for v in required) or len(set(required))!=len(required): fail('Title, summary, sections and required section IDs required.')
    ids=[]
    for section in sections:
        if not isinstance(section,dict) or not text(section.get('id')) or not text(section.get('heading')) or not isinstance(section.get('paragraphs'),list) or any(not isinstance(v,str) for v in section['paragraphs']) or not isinstance(section.get('tables'),list): fail('Invalid section contract.')
        ids.append(section['id'])
        for table in section['tables']:
            if not isinstance(table,dict) or not isinstance(table.get('headers'),list) or not table['headers'] or any(not isinstance(v,str) for v in table['headers']) or not isinstance(table.get('rows'),list) or any(not isinstance(row,list) or len(row)!=len(table['headers']) or any(not isinstance(v,str) for v in row) for row in table['rows']): fail('Table header/row mismatch.')
    if len(set(ids))!=len(ids) or any(value not in ids for value in required): fail('Missing or duplicate required section.')
    document=Document(); document.core_properties.author='AI Workflow Studio'; document.add_heading(title,0); document.add_paragraph(summary); table_count=0
    for section in sections:
        document.add_heading(section['heading'],1)
        for paragraph in section['paragraphs']: document.add_paragraph(paragraph)
        for value in section['tables']:
            table=document.add_table(rows=1,cols=len(value['headers'])); table.style='Light Shading Accent 1'
            for cell,content in zip(table.rows[0].cells,value['headers']): cell.text=content
            for row in value['rows']:
                for cell,content in zip(table.add_row().cells,row): cell.text=content
            table_count+=1
    target=output_directory(data.get('outputDirectory'))
    with (target/'report.docx').open('xb') as stream: document.save(stream)
    reopened=Document(target/'report.docx')
    if len(reopened.tables)!=table_count: fail('Generated DOCX reopening mismatch.')
    return {'file':'report.docx','sections':ids,'tables':table_count,'inputSha256':digest(json_bytes({k:v for k,v in data.items() if k!='outputDirectory'})),'reopened':True}

try:
    raw=sys.stdin.buffer.read(1048577)
    if len(raw)>1048576: fail('Input exceeds 1 MiB.')
    data=json.loads(raw.decode('utf-8-sig'),object_pairs_hook=pairs,parse_constant=lambda value: fail('Nonfinite JSON number.'))
    if not isinstance(data,dict): fail('JSON object required.')
    result=run(data)
    sys.stdout.buffer.write(json_bytes(result))
except Exception as error:
    sys.stderr.buffer.write(('Operation failed: '+str(error)+'\n').encode('utf-8'))
    sys.exit(1)
